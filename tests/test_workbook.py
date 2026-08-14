from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from result_scraper.scraper import DrawResult
from result_scraper.workbook import HEADERS, SHEET_NAME, load_existing_results, upsert_results


class WorkbookTests(unittest.TestCase):
    def test_upsert_is_sorted_duplicate_safe_and_digit_formula_driven(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            upsert_results(
                output,
                [
                    DrawResult(date(2026, 8, 11), "33775", "https://www.thinhnam.net.in/a"),
                    DrawResult(date(2026, 8, 9), "01234", "https://www.thinhnam.net.in/b"),
                ],
            )
            upsert_results(
                output,
                [DrawResult(date(2026, 8, 11), "12345", "https://www.thinhnam.net.in/c")],
            )

            stored = load_existing_results(output)
            self.assertEqual(list(stored), [date(2026, 8, 9), date(2026, 8, 11)])
            self.assertEqual(stored[date(2026, 8, 9)].grand_prize, "01234")
            self.assertEqual(stored[date(2026, 8, 11)].grand_prize, "12345")

            workbook = load_workbook(output, data_only=False)
            sheet = workbook[SHEET_NAME]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
            self.assertEqual(sheet["B2"].value, 1234)
            self.assertEqual(sheet["B2"].number_format, "00000")
            self.assertEqual(sheet["C2"].value, "=INT(B2/10000)")
            self.assertEqual(sheet["D2"].value, "=INT(MOD(B2,10000)/1000)")
            self.assertEqual(sheet["E2"].value, "=INT(MOD(B2,1000)/100)")
            self.assertEqual(sheet["F2"].value, "=INT(MOD(B2,100)/10)")
            self.assertEqual(sheet["G2"].value, "=MOD(B2,10)")
            for column in "CDEFG":
                self.assertEqual(sheet[f"{column}2"].number_format, "0")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(len(sheet.tables), 1)
            self.assertEqual(next(iter(sheet.tables.values())).ref, "A1:G3")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
